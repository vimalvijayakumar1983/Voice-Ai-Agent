"""Bounded CSV/XLSX import; formulas, macros and remote links are never executed."""

import csv
from io import BytesIO, StringIO
from zipfile import BadZipFile, ZipFile

from pydantic import ValidationError

from app.schemas.dialer import CustomerInput

MAX_BYTES = 2_000_000
MAX_ROWS = 5000


def parse_customers(filename, content):
    if len(content) > MAX_BYTES:
        raise ValueError("File must be at most 2 MB")
    if filename.lower().endswith(".csv"):
        rows = csv.reader(StringIO(content.decode("utf-8-sig")))
        workbook = None
    elif filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 500 or sum(entry.file_size for entry in entries) > 10_000_000:
                    raise ValueError("Workbook expands beyond the safe import limit")
                if any("vbaProject" in entry.filename for entry in entries):
                    raise ValueError("Macro workbooks are unsupported")
            workbook = load_workbook(
                BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            if workbook.active is None:
                workbook.close()
                raise ValueError("Workbook has no active worksheet")
            if (workbook.active.max_column or 0) > len(CustomerInput.model_fields) or (
                workbook.active.max_row or 0
            ) > MAX_ROWS + 1:
                workbook.close()
                raise ValueError("Worksheet dimensions exceed the import limit")
            rows = workbook.active.iter_rows(values_only=True)
        except (BadZipFile, KeyError) as exc:
            raise ValueError("Invalid XLSX workbook") from exc
    else:
        raise ValueError("Upload a UTF-8 CSV or XLSX file")
    try:
        header = [str(value or "").strip() for value in next(rows, ())]
        if len(header) != len(set(header)) or not {"name", "phone_number", "company"} <= set(
            header
        ):
            raise ValueError("Unique headers must include name, phone_number and company")
        if set(header) - CustomerInput.model_fields.keys():
            raise ValueError("Unknown column: use the documented customer import template")
        result, errors, seen = [], [], set()
        count = 0
        for count, row in enumerate(rows, start=1):
            if count > MAX_ROWS:
                raise ValueError("A file may contain at most 5,000 customer rows")
            if not any(value not in (None, "") for value in row):
                continue
            if len(row) > len(header) and any(
                value not in (None, "") for value in row[len(header) :]
            ):
                errors.append({"row": count + 1, "error": "More values than column headers"})
                continue
            values = {key: value for key, value in zip(header, row) if value not in (None, "")}
            try:
                if any(
                    isinstance(value, str) and value.startswith("=") for value in values.values()
                ):
                    raise ValueError("Formulas are not accepted; provide values only")
                customer = CustomerInput.model_validate(values)
                key = (customer.company, customer.phone_number)
                if key in seen:
                    raise ValueError("Duplicate company and phone number in file")
                seen.add(key)
                result.append(customer)
            except (ValidationError, ValueError) as exc:
                message = "Invalid customer fields or missing consent reference"
                if not isinstance(exc, ValidationError):
                    message = str(exc)
                errors.append({"row": count + 1, "error": message})
        return result, errors[:100], count
    finally:
        if workbook is not None:
            workbook.close()
